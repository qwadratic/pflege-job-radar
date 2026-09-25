"""Bavaria Reha/Vorsorge facilities (RHV_2024 sheet) -> candidate clinics registry rows.

TASK-145's AC#1: employers behind a group's board are not limited to Krankenhausplan sites (Reha
clinics, Vorsorgeeinrichtungen never appear there at all) -- confirmed live 2026-09-24 that this is a
pure sourcing gap, not a classification one: the employer-name classifier already recognizes
Reha-shaped names as clinic (patterns.json employer.clinic.kurklinik: "reha-?zentrum|rehazentrum|
rehabilitationszentrum|reha-?fachklinik|sanatorium|kurklinik"), but a DB check found zero postings
ever attached to any employer with "reha" in its name -- no source has ever crawled one.

Source: "Verzeichnis der Krankenhäuser und Vorsorge- oder Rehabilitationseinrichtungen", Statistische
Ämter des Bundes und der Länder, Stand 31.12.2024 (statistikportal.de/de/veroeffentlichungen/
krankenhausverzeichnis). Unlike the Krankenhausplan PDF, this is a structured export (name/town/PLZ/
operator/beds/fachgebiete each their own column) -- no OCR, no positional name/town ambiguity, so
there is no equivalent of that script's parse_quality uncertainty tiers.

Modeling decision (the "clinics vs employers" fork TASK-145 flagged, resolved here): reuse the
existing `clinics` table/CLINIC_SPEC as-is instead of remodeling into a renamed "employers" table.
Nothing downstream branches on Krankenhaus vs Reha -- Matcher, career_crawl, the vendor adapters, and
the clinic_links edge op all key purely off clinic_id/careers_url/ats_type/town/operator. Renaming the
table would touch every file in the pipeline for zero functional gain. The one real discriminator
needed (so Reha rows are visible as such, and so app/targets.py's only status check --
`status != "nicht_mehr_im_plan"` in app/targets.py:clinics_for -- keeps including them) is the
`status` field, set here to the literal "Reha-Einrichtung" (a value distinct from every real
Krankenhausplan status, so it can never collide with or be mistaken for a KH-plan status).

clinic_id: Krankenhausplan KeZ are bare 5-digit numbers. To guarantee zero collision with that space
forever, and make a Reha row visually identifiable in logs/URLs, this mints "RH<RH_ID_Pseudo>" from
the registry's own pseudonymous per-facility ID. app/autopilot/seed.py's CSV fallback path filters
clinic_id via .isdigit() -- an "RH..." id would be silently dropped there; that call site's own fix
was fixed directly (TASK-148 AC#4, app/autopilot/seed.py: the isdigit() check now also accepts
"RH<digits>"). The actual career-page discovery per facility (TASK-148 AC#2) is a separate,
cost-gated follow-up step, run only after these rows exist live -- crawlers/ats_discover2.py's free
homepage/sitemap/bewerben-button fingerprinting only ever probes clinics already IN the live registry.

  python data/sync_rhv_reha.py --dry-run     # print counts + a sample, write nothing
  python data/sync_rhv_reha.py               # write data/registry/reha_bavaria.csv (CSV only, no push)
  python data/sync_rhv_reha.py --push        # push reha_bavaria.csv rows to Supabase via the ingest edge function
"""
import argparse
import csv
import json
import os
import sys

import openpyxl
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from pflege_jobs.schema import CLINIC_SPEC  # noqa: E402

XLSX = "data/registry/krankenhausverzeichnis_24.xlsx"
OUT_CSV = "data/registry/reha_bavaria.csv"
LAND_BAYERN = "09"
TRAEGERART = {"1": "oeffentlich", "2": "freigemeinnuetzig", "3": "privat"}
BEZIRKE = ["Oberbayern", "Niederbayern", "Oberpfalz", "Oberfranken", "Mittelfranken", "Unterfranken", "Schwaben"]
SOURCE = ("RHV 2024 (Verzeichnis der Vorsorge- oder Rehabilitationseinrichtungen, "
          "Statistische Aemter des Bundes und der Laender, Stand 31.12.2024)")


def _fachgebiete(header, row, fa_labels):
    """FA_* columns hold a bed count per specialty; >0 means the facility runs that department."""
    out = []
    for col, label in fa_labels.items():
        i = header.index(col)
        v = row[i]
        if v and str(v).strip() not in ("", "0"):
            out.append(label)
    return ", ".join(out)


def _fa_labels(wb):
    ws = wb["DSB_RHV_2024"]
    labels = {}
    for r in ws.iter_rows(values_only=True):
        if r[0] and str(r[0]).startswith("FA_"):
            labels[r[0]] = r[1]
    return labels


def _kreis_names(wb):
    """'09173' -> 'Bad Toelz-Wolfratshausen'. RHV's own Kreis column is only the 3-digit suffix."""
    ws = wb["Kreis"]
    out = {}
    for r in ws.iter_rows(values_only=True):
        if r and r[0] and str(r[0]).startswith(LAND_BAYERN):
            out[str(r[0])] = r[2]
    return out


def _website(v):
    v = (v or "").strip()
    if not v:
        return ""
    return v if v.startswith("http") else "https://" + v


def parse(xlsx_path=XLSX):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["RHV_2024"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[2])
    fa_labels = _fa_labels(wb)
    kreis_names = _kreis_names(wb)
    idx = {k: i for i, k in enumerate(header)}
    out = []
    for r in rows[3:]:
        if r[idx["Land"]] != LAND_BAYERN:
            continue
        kreis3 = str(r[idx["Kreis"]])
        kreis_code = LAND_BAYERN + kreis3
        row = {k: "" for k, _ in CLINIC_SPEC}
        row["clinic_id"] = "RH" + str(r[idx["RH_ID_Pseudo"]])
        row["name"] = r[idx["RH_Name"]]
        row["town"] = r[idx["Ort"]]
        row["operator"] = r[idx["Trägername"]]
        row["traegerart"] = TRAEGERART.get(str(r[idx["Trägerart"]]), "")
        row["landkreis"] = kreis_names.get(kreis_code, "")
        row["regierungsbezirk"] = BEZIRKE[int(kreis3[0]) - 1]
        row["beds"] = r[idx["Betten"]] or 0
        row["fachrichtungen"] = _fachgebiete(header, r, fa_labels)
        row["website"] = _website(r[idx["Internet"]])
        row["status"] = "Reha-Einrichtung"
        row["parse_quality"] = "ok"
        row["source"] = SOURCE
        row["careers_url"] = ""
        row["ats_type"] = ""
        row["day_places"] = ""
        out.append(row)
    return out


INT_COLS = {"beds", "day_places"}


def _for_push(row):
    """CSV round-tripping turns Python None into "" (csv has no null); the edge function's
    json_to_recordset casts beds/day_places straight to Postgres int, and int4in("") raises
    'invalid input syntax for type integer: \"\"' -- confirmed live -- which fails the whole
    batch's single INSERT...SELECT statement, so every row in that batch is silently dropped."""
    return {k: (None if k in INT_COLS and v == "" else v) for k, v in row.items()}


def push(rows):
    """Same ingest edge-function op data/sync_krankenhausplan_2026.py uses for clinics rows."""
    rows = [_for_push(r) for r in rows]
    ing = {"Authorization": "Bearer " + os.environ["SUPABASE_ANON_KEY"], "apikey": os.environ["SUPABASE_ANON_KEY"],
           "x-ingest-secret": os.environ["PFLEGE_INGEST_SECRET"], "Content-Type": "application/json"}
    n = 0
    for i in range(0, len(rows), 50):
        r = requests.post(os.environ["PFLEGE_INGEST_URL"], headers=ing,
                          data=json.dumps({"clinics": rows[i:i + 50]}, ensure_ascii=False).encode("utf-8"), timeout=180)
        if r.status_code != 200:
            print("  push failed", r.status_code, r.text[:200]); continue
        n += r.json().get("clinics", 0) or 0
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--push", action="store_true", help="push data/registry/reha_bavaria.csv rows to Supabase")
    a = ap.parse_args()

    if a.push:
        rows = list(csv.DictReader(open(OUT_CSV, encoding="utf-8")))
        n = push(rows)
        print(f"pushed {n}/{len(rows)} rows to Supabase clinics")
        return

    rows = parse()
    no_website = [r["clinic_id"] for r in rows if not r["website"]]
    print(f"Bavaria RHV facilities: {len(rows)}")
    print(f"  no website          : {len(no_website)}  {no_website}")
    print(f"  bed range           : {min(r['beds'] for r in rows)}-{max(r['beds'] for r in rows)}")
    if a.dry_run:
        print("\n--dry-run: nothing written; sample:")
        for r in rows[:5]:
            print(f"  {r['clinic_id']:<8} {r['name'][:44]:<44} {r['town']:<20} beds={r['beds']:<4} {r['website']}")
        return
    with open(OUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _ in CLINIC_SPEC])
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {OUT_CSV} ({len(rows)} rows) -- not pushed to Supabase")


if __name__ == "__main__":
    main()
